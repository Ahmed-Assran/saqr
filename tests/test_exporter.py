"""
Tests for Lead Exporter Engine (OpenPyXL Excel & UTF-8 BOM CSV)
"""

from pathlib import Path
import openpyxl
import pytest

from src.stage2_scrapers.models import NormalizedLead, ScoredLead, TriggerEvent
from src.stage3_scorer.exporter import CATEGORY_TAB_NAMES, LeadExporter


def make_scored_lead(
    user_id: str,
    username: str,
    category: str,
    total_score: float,
) -> ScoredLead:
    trigger = TriggerEvent(
        platform="x",
        trigger_type="search_tweet",
        content=f"تجربة تداول في {category}",
        url=f"https://x.com/{username}/status/100",
    )
    lead = NormalizedLead(
        platform="x",
        user_id=user_id,
        username=username,
        display_name=f"User {username}",
        bio="مستثمر في السوق الأمريكي 🇸🇦",
        location="الرياض",
        profile_url=f"https://x.com/{username}",
        can_dm=True,
        triggers=[trigger],
    )
    return ScoredLead(
        lead=lead,
        heuristic_score=total_score * 0.6,
        llm_intent_score=total_score * 0.4,
        total_score=total_score,
        is_genuine_retail_investor=True,
        is_competitor_or_seller=False,
        interest_category=category,
        qualification_reason_ar="مستثمر نشط ومؤهل",
    )


def test_select_top_leads_balanced(tmp_path):
    exporter = LeadExporter(exports_dir=tmp_path, target_count=10, min_score=55.0)

    # Create 20 leads across different categories
    leads = []
    categories = list(CATEGORY_TAB_NAMES.keys())
    for i in range(20):
        cat = categories[i % len(categories)]
        leads.append(make_scored_lead(f"u_{i}", f"user_{i}", cat, 60.0 + i))

    selected = exporter.select_top_leads(leads)
    assert len(selected) == 10
    # Confirm multiple categories are represented
    selected_cats = {sl.interest_category for sl in selected}
    assert len(selected_cats) >= 4


def test_export_excel_and_csv_files(tmp_path):
    exporter = LeadExporter(exports_dir=tmp_path, target_count=50, min_score=55.0)

    leads = [
        make_scored_lead("1", "opt_lead", "1_Options_Traders", 85.0),
        make_scored_lead("2", "sharia_lead", "2_Halal_Sharia_Stocks", 80.0),
        make_scored_lead("3", "swing_lead", "3_Active_Day_Swing_Traders", 75.0),
        make_scored_lead("4", "tech_lead", "4_Specific_Tech_Tickers", 70.0),
        make_scored_lead("5", "broker_lead", "5_Broker_Platform_Seekers", 65.0),
    ]

    res = exporter.export_all(leads, date_tag="test_2026")
    excel_path = res["excel"]
    csv_path = res["csv"]

    assert excel_path.exists()
    assert csv_path.exists()
    assert excel_path.stat().st_size > 0
    assert csv_path.stat().st_size > 0

    # Inspect Excel Workbook
    wb = openpyxl.load_workbook(excel_path)
    sheet_names = wb.sheetnames
    assert "All_Top_50" in sheet_names
    for tab in CATEGORY_TAB_NAMES.values():
        assert tab in sheet_names

    ws_all = wb["All_Top_50"]
    assert ws_all.views.sheetView[0].rightToLeft is True
    assert ws_all.freeze_panes == "A2"
    # Row 1 is header, Row 2-6 are leads
    assert ws_all.max_row == 6

    # Verify contact columns exist in headers
    headers = [cell.value for cell in ws_all[1]]
    assert "رقم الجوال" in headers
    assert "رابط واتساب" in headers
    assert "حساب تيليجرام" in headers

    # Inspect CSV BOM
    with open(csv_path, "rb") as f:
        bom = f.read(3)
        assert bom == b"\xef\xbb\xbf"  # UTF-8 with BOM
